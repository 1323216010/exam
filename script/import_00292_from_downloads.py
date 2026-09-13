#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Import 市政学 papers from the local downloads folder into json/00292."""

from __future__ import annotations

import json
import os
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "json" / "00292"
EXTRACT_DIR = ROOT / ".tmp_extract"
DOWNLOADS = Path(r"E:\Downloads")

TYPE_MAP = {
    "单选题": "单项选择题",
    "单项选择题": "单项选择题",
    "多选题": "多项选择题",
    "多项选择题": "多项选择题",
    "简答题": "简答题",
    "论述题": "论述题",
    "材料分析题": "材料分析题",
    "材料题": "材料分析题",
    "名词解释": "名词解释",
    "名词解释题": "名词解释",
}


def source_dir() -> Path:
    for item in DOWNLOADS.iterdir():
        if item.is_dir() and item.name.startswith("00292"):
            return item
    raise FileNotFoundError("00292 downloads folder not found")


def extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append("\t".join(cells))
    return "\n".join(parts)


def extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


_WORD = None


def get_word():
    global _WORD
    if _WORD is None:
        import win32com.client

        _WORD = win32com.client.Dispatch("Word.Application")
        _WORD.Visible = False
        _WORD.DisplayAlerts = 0
    return _WORD


def quit_word():
    global _WORD
    if _WORD is not None:
        try:
            _WORD.Quit()
        except Exception:
            pass
        _WORD = None


def extract_doc(path: Path) -> str:
    word = get_word()
    doc = word.Documents.Open(str(path), ReadOnly=True)
    try:
        return doc.Content.Text.replace("\r", "\n")
    finally:
        doc.Close(False)


def extract_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".doc":
        return extract_doc(path)
    return ""


def parse_date(name: str) -> tuple[str, str]:
    match = re.search(r"(20\d{2})\s*年?\s*(1[0-2]|0?[1-9])", name)
    if not match:
        return "", name
    year = match.group(1)
    month = str(int(match.group(2)))
    return f"{year}年{month}月", f"{year}年{month}月自考00292市政学试题及答案.json"


def normalize_type(raw: str, number: int | None = None) -> str:
    raw = (raw or "").strip()
    if raw in TYPE_MAP:
        return TYPE_MAP[raw]
    if "判断" in raw:
        return "判断改错题"
    if number is not None:
        if number <= 25:
            return "单项选择题"
        if number <= 33:
            return "多项选择题"
        if number <= 39:
            return "简答题"
        if number == 40:
            return "论述题"
        if number >= 41:
            return "材料分析题"
    return raw or "简答题"


def clean_answer(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^正确答案[：:]\s*", "", text)
    text = re.sub(r"^答案[：:]\s*", "", text)
    text = re.sub(r"^答[：:]\s*", "", text)
    text = re.sub(r"教材\s*P\S*\s*", "", text)
    text = re.sub(r"^【|】$", "", text)
    return text.strip()


def strip_section_tail(text: str) -> str:
    return re.sub(r"[一二三四五六七八九十]+、\S.*$", "", text).strip()


def parse_format_b(text: str) -> list[dict]:
    pattern = re.compile(
        r"(?m)^(\d{1,2})[、.．]\s*【([^】]+)】\s*(.*?)(?=^\d{1,2}[、.．]\s*【|\Z)",
        re.S,
    )
    questions = []
    for match in pattern.finditer(text):
        number = int(match.group(1))
        qtype = normalize_type(match.group(2), number)
        body = match.group(3).strip()
        answer_match = re.search(r"(?:^|\n)\s*(?:正确答案|答案)\s*[：:]\s*(.*)", body, re.S)
        answer = ""
        if answer_match:
            answer = clean_answer(answer_match.group(1).split("解析")[0])
            body = body[: answer_match.start()].strip()
        options = {}
        for opt in re.finditer(r"(?m)^([A-E])[:：.．、]\s*(.+)$", body):
            options[opt.group(1)] = opt.group(2).strip()
        content = re.sub(r"(?m)^[A-E][:：.．、]\s*.+$", "", body).strip()
        content = re.sub(r"\n{2,}", "\n", content)
        item = {
            "question_number": str(number),
            "question_type": qtype,
            "content": content,
            "answer": answer,
            "score": None,
        }
        if options:
            item["options"] = options
        questions.append(item)
    return questions


SECTION_TYPES = [
    (re.compile(r"单项选择"), "单项选择题"),
    (re.compile(r"多项选择"), "多项选择题"),
    (re.compile(r"判断"), "判断改错题"),
    (re.compile(r"名词解释"), "名词解释"),
    (re.compile(r"简答"), "简答题"),
    (re.compile(r"论述"), "论述题"),
    (re.compile(r"材料|案例分析"), "材料分析题"),
]


def infer_type_from_sections(text: str, number: int) -> str:
    current = ""
    for line in text.splitlines():
        for pattern, qtype in SECTION_TYPES:
            if re.match(r"^[一二三四五六七八九十]、", line) and pattern.search(line):
                current = qtype
                break
        if re.match(rf"^{number}[、.．]", line.strip()) and current:
            return current
    return normalize_type("", number)


def extract_inline_choice_answer(text: str) -> str:
    match = re.search(r"[\(（]\s*([A-E]{1,5})\s*[\)）]", text)
    if match:
        return match.group(1)
    match = re.search(r"【\s*([A-E]{1,5}|[√×])\s*】", text)
    if match:
        return match.group(1)
    return ""


def apply_judgment_table(text: str, questions: list[dict]) -> None:
    by_number = {q["question_number"]: q for q in questions}
    for match in re.finditer(r"(?m)^(\d{1,2})\t([√×对错])\t?(.*)$", text):
        number, mark, correction = match.group(1), match.group(2), match.group(3).strip()
        item = by_number.get(number)
        if not item:
            continue
        item["question_type"] = "判断改错题"
        item["answer"] = f"{mark} {correction}".strip()


def parse_format_a(text: str) -> list[dict]:
    text = re.sub(r"(（\d+分）)\s*(\d{1,2}[、.．])", r"\1\n\2", text)
    pattern = re.compile(
        r"(?m)^(\d{1,2})[、.．]\s*(?:\d+[、.．]\s*)?(.*?)(?=^\d{1,2}[、.．]\s*|\Z)",
        re.S,
    )
    skip_starts = ("答题前", "每小题", "本大题", "考生务必", "将所有试题", "判断下列")
    questions = []
    for match in pattern.finditer(text):
        number = int(match.group(1))
        if number < 1 or number > 60:
            continue
        body = match.group(2).strip()
        if any(body.startswith(prefix) for prefix in skip_starts):
            continue
        qtype = infer_type_from_sections(text, number)
        answer = ""
        answer_match = re.search(r"(?:正确答案|参考答案|答案)\s*[：:]\s*(.*)", body, re.S)
        if answer_match:
            raw_answer = answer_match.group(1).strip()
            if qtype.endswith("选择题"):
                letter = re.match(r"[A-E]{1,5}", raw_answer.replace(" ", "").replace("、", ""))
                answer = letter.group(0) if letter else clean_answer(raw_answer.split("（")[0])
            else:
                answer = clean_answer(raw_answer)
                answer = re.sub(r"\n{2,}", "\n", answer).strip()
            body = body[: answer_match.start()].strip()
        else:
            da_match = re.search(r"(?:^|\n)\s*答[：:]\s*(.*)", body, re.S)
            if da_match:
                answer = clean_answer(da_match.group(1))
                body = body[: da_match.start()].strip()
            else:
                answer = extract_inline_choice_answer(body)
        options = {}
        for opt in re.finditer(r"([A-E])[.．、]\s*([^\t\n]+)", body):
            options[opt.group(1)] = opt.group(2).strip()
        content = re.sub(r"[A-E][.．、]\s*[^\t\n]+", "", body)
        content = re.sub(r"[\(（]\s*[A-E]{1,5}\s*[\)）]", "", content)
        content = content.replace("【", "").replace("】", "")
        content = strip_section_tail(content)
        content = re.sub(r"\s+", " ", content).strip()
        if not content or len(content) < 4:
            continue
        item = {
            "question_number": str(number),
            "question_type": qtype,
            "content": content,
            "answer": answer,
            "score": None,
        }
        if options:
            item["options"] = options
        questions.append(item)
    apply_judgment_table(text, questions)
    return questions


def parse_answer_key(text: str) -> dict[str, str]:
    answers: dict[str, str] = {}
    for match in re.finditer(r"(\d{1,2})\s*[．.、]\s*([A-E]{1,5})\b", text):
        answers[match.group(1)] = match.group(2)
    for match in re.finditer(
        r"(?m)^(\d{2})\s*[．.、]\s*(.+?)(?=^\d{2}\s*[．.、]|\Z)",
        text,
        re.S,
    ):
        number = match.group(1)
        if int(number) < 34:
            continue
        body = match.group(2).strip()
        if re.fullmatch(r"[A-E]{1,5}", body.split()[0] if body else ""):
            continue
        answers[number] = re.sub(r"\n{2,}", "\n", body).strip()
    return answers


def apply_answer_key(questions: list[dict], answers: dict[str, str]) -> None:
    for item in questions:
        number = item["question_number"]
        if answers.get(number) and not str(item.get("answer") or "").strip():
            item["answer"] = answers[number]


def parse_ocr_markdown(text: str) -> list[dict]:
    text = re.sub(r"注意事项：.*?(?=\n#{0,4}\s*[一二三四五]|\n\d{1,2}[.、．]\s*\*?[^\n]{8,})", "", text, flags=re.S)
    pattern = re.compile(
        r"(?m)^\*{0,2}\s*(\d{1,2})[.、．]\s*(?:\*{0,2})(.*?)(?=^\*{0,2}\s*\d{1,2}[.、．]\s*|\Z)",
        re.S,
    )
    skip_starts = ("本试卷", "应考者", "涂写", "答题前", "每小题", "将所有", "判断下列")
    questions = []
    for match in pattern.finditer(text):
        number = int(match.group(1))
        if number < 1 or number > 60:
            continue
        body = match.group(2).strip()
        body = body.replace("**", "")
        if any(body.startswith(prefix) for prefix in skip_starts):
            continue
        if re.fullmatch(r"[A-E]{1,5}", body.replace(" ", "").replace(",", "").replace("、", "")[:8] or ""):
            continue
        qtype = infer_type_from_sections(text, number)
        answer = ""
        answer_match = re.search(r"答案\s*[：:]\s*(.*)", body, re.S)
        if answer_match:
            raw = answer_match.group(1).strip()
            if qtype.endswith("选择题") or number <= 33:
                letters = re.sub(r"[^A-E]", "", raw.split("\n")[0])
                answer = letters or clean_answer(raw)
            else:
                answer = clean_answer(raw)
            body = body[: answer_match.start()].strip()
        options = {}
        for opt in re.finditer(r"(?m)^\s*-?\s*([A-E])[:：.．、]\s*(.+)$", body):
            options[opt.group(1)] = opt.group(2).strip()
        content = re.sub(r"(?m)^\s*-?\s*[A-E][:：.．、]\s*.+$", "", body)
        content = strip_section_tail(content)
        content = re.sub(r"\s+", " ", content).strip(" -*")
        if not content or len(content) < 6:
            continue
        if options and number <= 25:
            qtype = "单项选择题"
        elif options and number <= 33:
            qtype = "多项选择题"
        item = {
            "question_number": str(number),
            "question_type": qtype,
            "content": content,
            "answer": answer,
            "score": None,
        }
        if options:
            item["options"] = options
        questions.append(item)
    return questions


def parse_questions(text: str) -> list[dict]:
    if re.search(r"^\s*-?\s*[A-E][:：]\s+", text, re.M) or "**答案" in text:
        questions = parse_ocr_markdown(text)
        if len(questions) >= 20:
            pass
        elif "【单选题】" in text or "【多选题】" in text:
            questions = parse_format_b(text)
        else:
            questions = parse_format_a(text)
    elif "【单选题】" in text or "【多选题】" in text:
        questions = parse_format_b(text)
    else:
        questions = parse_format_a(text)
    apply_answer_key(questions, parse_answer_key(text))
    unique: dict[str, dict] = {}
    for item in questions:
        key = item["question_number"]
        old = unique.get(key)
        if old is None:
            unique[key] = item
            continue
        if item.get("options") and not old.get("options"):
            if old.get("answer") and not item.get("answer"):
                item["answer"] = old["answer"]
            unique[key] = item
        elif str(item.get("answer") or "").strip() and not str(old.get("answer") or "").strip():
            old["answer"] = item["answer"]
    result = list(unique.values())
    result.sort(key=lambda q: int(re.sub(r"\D", "", q["question_number"]) or 0))
    return result


def ai_convert(text: str, date_label: str) -> dict | None:
    from openai import OpenAI

    client = OpenAI(
        api_key=os.getenv("DASHSCOPE_API_KEY") or os.getenv("YINLI_API_KEY"),
        base_url="https://dashscope.aliyuncs.com/compatible-mode/v1"
        if os.getenv("DASHSCOPE_API_KEY")
        else "https://yinli.one/v1",
    )
    model = "qwen-plus" if os.getenv("DASHSCOPE_API_KEY") else "gemini-3-flash-preview"
    prompt = f"""把下面这份00292市政学试卷转成JSON。必须包含全部题目和能找到的答案。
日期用：{date_label}
科目：市政学
课程代码：00292
格式：
{{
  "exam_info": {{"title": "...", "subject": "市政学", "code": "00292", "date": "{date_label}"}},
  "questions": [
    {{"question_number":"1","question_type":"单项选择题","content":"...","options":{{"A":"...","B":"...","C":"...","D":"..."}},"answer":"A","score":null}}
  ]
}}
选择题答案只保留字母。主观题答案保留原文要点。不要编造找不到的答案，找不到就填空字符串。

试卷原文：
{text[:24000]}
"""
    completion = client.chat.completions.create(
        model=model,
        messages=[{"role": "user", "content": prompt}],
        response_format={"type": "json_object"},
    )
    raw = completion.choices[0].message.content or ""
    start, end = raw.find("{"), raw.rfind("}")
    if start == -1 or end == -1:
        return None
    return json.loads(raw[start : end + 1])


def quality_ok(questions: list[dict]) -> bool:
    return len(questions) >= 20


def build_exam(date_label: str, questions: list[dict]) -> dict:
    from fill_00292_scores import parse_date, scores_for

    year, month = parse_date(f"{date_label}自考")
    table = scores_for(year, month, [q.get("question_type") or "" for q in questions])
    for item in questions:
        if item.get("score") in (None, ""):
            item["score"] = table.get(item.get("question_type") or "", 1)
    return {
        "exam_info": {
            "title": f"{date_label}高等教育自学考试市政学试题",
            "subject": "市政学",
            "code": "00292",
            "date": date_label,
        },
        "questions": questions,
    }


def log(msg: str) -> None:
    print(msg, flush=True)
    log_path = ROOT / ".tmp_import_log.txt"
    with log_path.open("a", encoding="utf-8") as fh:
        fh.write(msg + "\n")


def main() -> None:
    (ROOT / ".tmp_import_log.txt").write_text("", encoding="utf-8")
    src = source_dir()
    EXTRACT_DIR.mkdir(exist_ok=True)
    OUT_DIR.mkdir(exist_ok=True)

    files = [f for f in sorted(src.iterdir()) if f.suffix.lower() in {".doc", ".docx", ".pdf"}]
    log(f"source={src}")
    log(f"files={len(files)}")

    results = []
    for index, path in enumerate(files, 1):
        date_label, out_name = parse_date(path.name)
        if not date_label:
            log(f"[{index}/{len(files)}] skip unnamed {path.name}")
            continue
        log(f"[{index}/{len(files)}] {date_label} {path.suffix} {path.stat().st_size}")
        try:
            text = extract_file(path)
        except Exception as exc:
            log(f"  extract failed: {exc}")
            text = ""
        extract_path = EXTRACT_DIR / f"{date_label}{path.suffix}.txt"
        extract_path.write_text(text, encoding="utf-8")
        questions = parse_questions(text) if text.strip() else []
        answers = parse_answer_key(text) if text.strip() else {}
        if not quality_ok(questions) and answers:
            existing = OUT_DIR / out_name
            old_name = out_name.replace("试题及答案", "试题")
            candidates = [existing, OUT_DIR / old_name]
            for cand in candidates:
                if not cand.exists():
                    continue
                try:
                    old = json.loads(cand.read_text(encoding="utf-8"))
                except Exception:
                    continue
                old_qs = old.get("questions") or []
                if quality_ok(old_qs):
                    apply_answer_key(old_qs, answers)
                    questions = old_qs
                    log(f"  merged answer key into {cand.name}")
                    break
        if not quality_ok(questions):
            if answers:
                key_path = EXTRACT_DIR / f"{date_label}.answers.json"
                key_path.write_text(json.dumps(answers, ensure_ascii=False, indent=2), encoding="utf-8")
                log(f"  answer-key only ({len(answers)} answers), no questions")
                results.append((date_label, out_name, len(answers), "answers-only"))
            else:
                log(f"  too little text ({len(text)}), skip")
                results.append((date_label, out_name, None, "scan"))
            continue
        data = build_exam(date_label, questions)

        data.setdefault("exam_info", {})
        data["exam_info"]["subject"] = "市政学"
        data["exam_info"]["code"] = "00292"
        data["exam_info"]["date"] = date_label
        out_path = OUT_DIR / out_name
        out_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        qn = len(data.get("questions") or [])
        ans = sum(1 for q in data["questions"] if str(q.get("answer") or "").strip())
        log(f"  saved {out_name} questions={qn} answers={ans} ok={quality_ok(questions)}")
        results.append((date_label, out_name, qn, "ok"))

    quit_word()
    log("\nDONE")
    for item in results:
        log(str(item))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    try:
        main()
    finally:
        quit_word()
